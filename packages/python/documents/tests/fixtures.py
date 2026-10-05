"""Builds realistic sample files of every supported format, in memory."""

from __future__ import annotations

import contextlib
import io
import os
import shutil
import tempfile

LONG_PARAGRAPH = " ".join(
    f"Sentence number {i} explains how the settlement batch reconciles ledger entries against the bank feed."
    for i in range(1, 60)
)


def make_docx(extra_intro: str | None = None) -> bytes:
    import docx
    from docx.shared import Pt  # noqa: F401

    d = docx.Document()
    d.core_properties.title = "Refund Operations Handbook"
    d.core_properties.author = "Payments Team"
    if extra_intro:
        d.add_paragraph(extra_intro)
    d.add_heading("Refund policy", level=1)
    d.add_paragraph(
        "Customers may request a refund within 30 days of purchase. Refunds are issued to the original payment method."
    )
    d.add_heading("Approval limits", level=2)
    d.add_paragraph(
        "Any refund above the agent limit requires approval from a team lead. See ticket BILL-1042 for the history."
    )
    t = d.add_table(rows=4, cols=3)
    for i, h in enumerate(["Role", "Region", "Max refund"]):
        t.rows[0].cells[i].text = h
    data = [["Agent", "US", "500"], ["Team lead", "US", "5000"], ["Agent", "EU", "400"]]
    for r, row in enumerate(data, start=1):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    merged = t.rows[3].cells[0].merge(t.rows[3].cells[1])  # horizontal merge -> must not duplicate
    merged.text = "Agent EU"
    d.add_paragraph("Escalate first", style="List Bullet")
    d.add_paragraph("Then document the case", style="List Bullet 2")
    code = d.add_paragraph()
    run = code.add_run("def issue_refund(order_id):")
    run.font.name = "Consolas"
    code2 = d.add_paragraph()
    run2 = code2.add_run("    return payments.refund(order_id)")
    run2.font.name = "Consolas"
    d.add_heading("Settlement", level=1)
    d.add_paragraph(LONG_PARAGRAPH)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows: int = 60) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["B2"] = "Q3 price list"
    ws.append([])
    ws.append([None, "SKU", "Product", "Region", "Price"])  # row 4, starts at column B
    regions = ["US", "EU", "APAC"]
    for i in range(rows):
        ws.append([None, f"SKU{10000 + i}", f"Widget model {i}", regions[i % 3], 10 + i * 1.5])
    ws.append([])
    ws.append([None, "Currency", "Rate"])
    ws.append([None, "EUR", 1.08])
    ws.append([None, "GBP", 1.27])
    hidden = wb.create_sheet("Secret")
    hidden["A1"] = "should not be indexed"
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MARKDOWN = """---
title: Billing runbook
tags: [billing, ops]
---
# Billing runbook

The billing service owns invoices and refunds.

## Refunds

Refunds over the limit need approval. Track them under BILL-1042 and call `issue_refund()`.

- Verify the order
  - Check fraud score
- Issue the refund

```python
def issue_refund(order_id):
    return api.refund(order_id)
```

| Region | Limit |
|---|---|
| US | 500 |
| EU | 400 |

## Escalation

Page the on-call engineer via PD-77.
"""


def make_markdown() -> bytes:
    return MARKDOWN.encode()


TEXT = """Incident notes
==============

The payment gateway returned error E4012 for about twenty minutes.
Café orders in Zürich were affected.

ROOT CAUSE

A certificate rotation on the gateway failed.

- Rotate the certificate
- Add an expiry alert
"""


def make_text_latin1() -> bytes:
    return TEXT.encode("latin-1")


CSV = "Employee,Team,Location\nAna,Billing,Lisbon\nBo,Payments,Oslo\n"


def make_csv() -> bytes:
    return CSV.encode()


def make_vsdx() -> bytes:
    import vsdx
    from vsdx import VisioFile
    from vsdx.connectors import Connect

    media = os.path.join(os.path.dirname(vsdx.__file__), "media", "media.vsdx")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "flow.vsdx")
        out = os.path.join(tmp, "flow_out.vsdx")
        shutil.copy(media, src)
        with contextlib.redirect_stdout(io.StringIO()), VisioFile(src) as vis:
            page = vis.pages[0]
            by = {s.ID: s for s in page.all_shapes}
            by["1"].text = "Receive order"
            by["2"].text = "Validate order"
            by["3"].text = "submit"
            by["5"].text = ""
            by["7"].text = "Ship order"
            by["8"].text = "SLA is 24h"
            conn = Connect.create(page, by["2"], by["7"])
            conn.text = "approved"
            vis.save_vsdx(out)
        with open(out, "rb") as fh:
            return fh.read()


def _wrap(text: str, width: float, font: str, size: float) -> list[str]:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines: list[str] = []
    line = ""
    for word in text.split():
        candidate = f"{line} {word}".strip()
        if line and stringWidth(candidate, font, size) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    return [*lines, line] if line else lines


def make_pdf() -> bytes:
    """Three pages: a titled, sectioned page with bullets and a ruled table; a
    paragraph carried over the page break and a code block; a two-column page.
    Every page has a running header and a page-number footer."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    width, height = letter
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setTitle("Microsoft Word - refunds_v3.docx")  # an export's junk title: ignored
    c.setAuthor("Payments Team")

    def chrome(page: int) -> None:
        c.setFont("Helvetica", 9)
        c.drawString(72, height - 40, "ACME Corp - Internal")
        c.drawCentredString(width / 2, 30, f"Page {page} of 3")

    def line(x: float, y: float, text: str, font: str = "Helvetica", size: float = 11) -> None:
        c.setFont(font, size)
        c.drawString(x, y, text)

    # page 1
    chrome(1)
    line(72, 700, "Refund Operations Guide", "Helvetica-Bold", 22)
    line(72, 660, "Refund policy", "Helvetica-Bold", 16)
    line(72, 635, "Customers may request a refund within 30 days of purchase. Refunds are issued")
    line(72, 621, "to the original payment method once the nightly recon-")
    line(72, 607, "ciliation job has matched the ledger entry.")
    line(72, 580, "Approval limits", "Helvetica-Bold", 13)
    line(72, 558, "Any refund above the agent limit requires approval from a team lead.")
    line(90, 536, "• Escalate first")
    line(108, 522, "• Then document the case in BILL-1042")
    line(72, 496, "Table 1: Refund limits")
    top, row_h, cols = 480, 18, [72, 222, 372, 522]
    rows = [
        ["Role", "Region", "Max refund"],
        ["Agent", "US", "500"],
        ["Team lead", "US", "5000"],
        ["Agent", "EU", "400"],
    ]
    for i in range(len(rows) + 1):
        c.line(cols[0], top - i * row_h, cols[-1], top - i * row_h)
    for x in cols:
        c.line(x, top, x, top - len(rows) * row_h)
    for r, row in enumerate(rows):
        for col, value in enumerate(row):
            line(cols[col] + 4, top - r * row_h - 13, value, size=10)
    line(72, 380, "Settlement batches run nightly and every batch writes its totals to the")
    c.showPage()

    # page 2
    chrome(2)
    line(72, 720, "general ledger before the bank feed is compared.")
    line(72, 690, "Settlement", "Helvetica-Bold", 16)
    line(72, 665, "Run the reconciliation script when a batch fails:")
    line(72, 645, "def reconcile(batch_id):", "Courier", 10)
    line(96, 633, "return ledger.match(batch_id)", "Courier", 10)
    c.showPage()

    # page 3: a full-width heading and intro, then two columns
    chrome(3)
    line(72, 720, "Appendix", "Helvetica-Bold", 16)
    line(72, 700, "This appendix lists the settlement codes and the escalation contacts for every batch.")
    left = _wrap("Left column text explains the settlement codes. " * 12, 215, "Helvetica", 11)
    right = _wrap("Right column text covers the escalation contacts. " * 12, 215, "Helvetica", 11)
    for i, text in enumerate(left):
        line(72, 675 - i * 14, text)
    for i, text in enumerate(right):
        line(320, 675 - i * 14, text)
    c.showPage()
    c.save()
    return buf.getvalue()


def make_pdf_outline() -> bytes:
    """Headings set like body text, named by the PDF's bookmarks."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setFont("Helvetica", 11)
    y = 720
    for title, level, body in [
        ("1 Overview", 0, "The gateway settles card payments every night."),
        ("1.1 Scope", 1, "Only card payments are in scope for this runbook."),
        ("2 Contacts", 0, "Page the payments on-call engineer for failed batches."),
    ]:
        key = title.replace(" ", "_")
        c.bookmarkHorizontal(key, 72, y + 12)
        c.addOutlineEntry(title.split(" ", 1)[1], key, level=level)
        c.drawString(72, y, title)
        c.drawString(72, y - 30, body)
        y -= 70
    c.showPage()
    c.save()
    return buf.getvalue()


def make_pdf_scan() -> bytes:
    """A page with no text layer, like a scan."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.rect(72, 72, 400, 600, fill=1)
    c.showPage()
    c.save()
    return buf.getvalue()


def make_pptx() -> bytes:
    """A title slide; bullets with notes; a table and a chart; an untitled slide
    with a text box and a group; and a hidden slide."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    prs.core_properties.author = "Payments Team"

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Refund Operations Review"
    slide.placeholders[1].text = "Q3 2026"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Approval limits"
    body = slide.placeholders[1].text_frame
    body.text = "Agents approve refunds up to 500"
    for text, level in [("Team leads approve up to 5000", 1), ("See ticket BILL-1042", 0)]:
        paragraph = body.add_paragraph()
        paragraph.text = text
        paragraph.level = level
    slide.notes_slide.notes_text_frame.text = "Mention that the EU limit is lower."

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Limits by region"
    table = slide.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(4), Inches(1.2)).table
    for r, row in enumerate([["Role", "Region", "Max refund"], ["Agent", "US", "500"], ["", "", "400"]]):
        for col, value in enumerate(row):
            table.cell(r, col).text = value
    table.cell(2, 0).merge(table.cell(2, 1))
    table.cell(2, 0).text = "Agent EU"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Refunds", (120, 95))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4), Inches(3), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Refunds per quarter"

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: no title
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Settlement runs nightly."
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(3), Inches(2), Inches(1)).text_frame.text = "Bank feed"
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "Ledger"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Draft ideas"
    slide._element.set("show", "0")

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf_clauses() -> bytes:
    """Contract-style text: numbered bold headings ending in a period, lettered
    clauses whose lines wrap flush with the letter, a bold sentence filling a
    line inside a paragraph, and a heading overprinted as "fake bold"."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    y = 720.0

    def line(text: str, font: str = "Helvetica", x: float = 40) -> None:
        nonlocal y
        c.setFont(font, 10)
        c.drawString(x, y, text)
        y -= 12

    line("1. General.", "Helvetica-Bold")
    line("A. The software, any documentation and any data accompanying this agreement are licensed, not sold,")
    line("to you by the vendor for use only under the terms of this agreement, which govern every update.")
    y -= 12
    line("B. The vendor may make future updates available; they may not include every existing feature.")
    y -= 12
    for dx in (0, 0.4):  # printed twice, a hair apart
        c.setFont("Helvetica", 10)
        c.drawString(40 + dx, y, "2. Updates.")
    y -= 12
    line("Updates install automatically when your device is idle, so you always run the latest release and")
    line("WHEN AN UPDATE CHANGES THE TERMS, THE NEW TERMS APPLY ONCE IT IS INSTALLED ON", "Helvetica-Bold")
    line("your device, unless you turn automatic updates off in the settings before it installs.")
    c.showPage()
    c.save()
    return buf.getvalue()


SCAN_LINES = [
    ("Refund Operations", 72),
    ("Customers may request a refund within 30 days of purchase.", 40),
    ("Refunds are issued to the original payment method.", 40),
]


def make_scan_image(dpi: int = 300):
    """A letter-size page image with printed text, as a scanner makes it."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("L", (int(8.5 * dpi), int(11 * dpi)), 255)
    draw = ImageDraw.Draw(image)
    y = 300
    for text, px in SCAN_LINES:
        draw.text((300, y), text, fill=0, font=ImageFont.load_default(size=px))
        y += int(px * 1.8)
    return image


def make_pdf_scanned(pages: int = 1, with_text_page: bool = False) -> bytes:
    """Pages that are only a scanned image (no text layer); optionally a
    normal text page first."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    width, height = letter
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    if with_text_page:
        c.setFont("Helvetica", 11)
        c.drawString(72, 700, "This cover page has a real text layer.")
        c.showPage()
    scan = ImageReader(make_scan_image())
    for _ in range(pages):
        c.drawImage(scan, 0, 0, width=width, height=height)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_pdf_borderless_table() -> bytes:
    """Prose around a table drawn without any rules: columns only by alignment,
    one cell wrapping onto a second line."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setFont("Helvetica", 11)
    c.drawString(72, 720, "Refund limits differ by role and region, as the table below shows.")
    rows = [
        ("Role", "Region", "Max refund"),
        ("Agent", "US", "500"),
        ("Team lead", "US", "5000"),
        ("Agent", "EU", "400 per order,"),
        ("", "", "2000 per day"),  # the cell above, wrapped
        ("Director", "Global", "No limit"),
    ]
    y = 690
    for role, region, limit in rows:
        c.setFont("Helvetica-Bold" if role == "Role" else "Helvetica", 11)
        for x, text in ((72, role), (240, region), (400, limit)):
            if text:
                c.drawString(x, y, text)
        y -= 16
    c.setFont("Helvetica", 11)
    c.drawString(72, y - 20, "Anything above these limits needs approval from finance.")
    c.showPage()
    c.save()
    return buf.getvalue()


def make_ole(stream: str) -> bytes:
    """Just enough of a legacy Office (OLE) file for sniffing: the header
    signature and a directory entry naming its main stream."""
    return b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504 + stream.encode("utf-16-le") + b"\x00" * 64


def make_odt() -> bytes:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("content.xml", "<office:document-content/>")
    return buf.getvalue()


# A flowchart with the syntax people write: front matter, a directive and a
# comment, shapes, labels between pipes and inside links, & fan-out, a
# subgraph an edge points at, and styling that says nothing.
MERMAID_FLOWCHART = """---
title: Order flow
---
%%{init: {"theme": "dark"}}%%
flowchart TD
    %% entry point
    A[Receive order] --> B{Valid?}
    B -- yes --> C([Ship order]) & D[(Orders DB)]
    B -->|no| E>Reject order]
    subgraph pay [Payments]
      P1["Charge card<br/>via **Stripe**"] -.-> P2((Refund))
    end
    C --> pay
    X ~~~ Y
    classDef red fill:#f00
    class A red
"""

MERMAID_SEQUENCE = """sequenceDiagram
    title Checkout
    actor U as Shopper
    participant W as Web app
    participant API
    U->>W: Click pay
    W->>+API: POST /orders
    alt card ok
      API-->>-W: 201 Created
    else declined
      API-->>W: 402 Payment Required
    end
    loop every 5s
      W-)API: poll status
    end
    Note over W,API: retries are idempotent
"""


def make_mermaid() -> bytes:
    return MERMAID_FLOWCHART.encode()
