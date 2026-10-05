"""Excel (.xlsx) -> IR using openpyxl in streaming (read-only) mode.

Each sheet is split into rectangular regions separated by blank rows; each
region becomes one TABLE element. A lone 1-2 cell row directly above a region
("Q3 Pricing") becomes that table's caption. Formulas yield their cached
values (``data_only=True``); files never opened in Excel may have none.
"""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from typing import Any

from forge_task_documents._util import column_letter
from forge_task_documents.errors import ParseError
from forge_task_documents.models import (
    Element,
    ElementKind,
    ParsedDocument,
    SourceFile,
    SourceLocation,
    TableData,
)
from forge_task_documents.parsers.base import cell_to_str, dedupe_headers, title_from_filename


class XlsxParser:
    name = "xlsx"
    version = "1"
    extensions = frozenset({"xlsx", "xlsm"})
    media_types = frozenset({"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})

    def __init__(
        self,
        *,
        include_hidden_sheets: bool = False,
        max_rows_per_sheet: int = 200_000,
        max_cols: int = 200,
    ) -> None:
        self.include_hidden_sheets = include_hidden_sheets
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_cols = max_cols

    def parse(self, source: SourceFile) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise ParseError("openpyxl is not installed") from exc
        try:
            wb = openpyxl.load_workbook(source.open(), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError(f"cannot open xlsx {source.filename!r}: {exc}") from exc

        elements: list[Element] = []
        sheets: list[str] = []
        try:
            for ws in wb.worksheets:
                if ws.sheet_state != "visible" and not self.include_hidden_sheets:
                    continue
                sheets.append(ws.title)
                rows = ws.iter_rows(max_col=self.max_cols, values_only=True)
                elements.extend(self._sheet_elements(ws.title, rows))
        finally:
            wb.close()

        title = (wb.properties.title or "").strip() if wb.properties else ""
        return ParsedDocument(
            tenant_id=source.tenant_id,
            doc_id=source.doc_id,
            source_type="xlsx",
            title=title or title_from_filename(source.filename),
            elements=elements,
            parser_name=self.name,
            parser_version=self.version,
            metadata={"sheets": sheets},
        )

    def _sheet_elements(self, sheet: str, rows: Iterable[tuple[Any, ...]]) -> Iterator[Element]:
        pending_caption: str | None = None
        table_no = 0
        for first_row, col_offset, region in _regions(rows, self.max_rows_per_sheet):
            non_empty = [c for c in region[0] if c]
            if len(region) == 1 and len(non_empty) <= 2:
                if pending_caption:  # two captions in a row: keep the first as text
                    yield _paragraph(pending_caption, sheet)
                pending_caption = " ".join(non_empty)
                continue
            table_no += 1
            yield _region_to_table(sheet, first_row, col_offset, region, pending_caption, table_no)
            pending_caption = None
        if pending_caption:
            yield _paragraph(pending_caption, sheet)


def _paragraph(text: str, sheet: str) -> Element:
    return Element(kind=ElementKind.PARAGRAPH, text=text, section_path=[sheet])


def _regions(rows: Iterable[tuple[Any, ...]], max_rows: int) -> Iterator[tuple[int, int, list[list[str]]]]:
    """Yield (1-based first row, 0-based column offset, trimmed rows) per block
    of non-blank rows."""
    current: list[list[str]] = []
    start = 0
    for r_idx, raw in enumerate(rows, start=1):
        if r_idx > max_rows:
            break
        cells = [cell_to_str(v) for v in raw]
        if not any(cells):
            if current:
                yield start, *_trim_columns(current)
                current = []
            continue
        if not current:
            start = r_idx
        current.append(cells)
    if current:
        yield start, *_trim_columns(current)


def _trim_columns(rows: list[list[str]]) -> tuple[int, list[list[str]]]:
    """Drop fully empty columns at both edges. Returns (left offset, rows)."""
    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]
    used = [i for i in range(width) if any(r[i] for r in padded)]
    lo, hi = used[0], used[-1]
    return lo, [r[lo : hi + 1] for r in padded]


def _looks_like_header(row: list[str]) -> bool:
    filled = [c for c in row if c]
    if len(filled) < max(1, len(row) // 2):
        return False
    numeric = sum(1 for c in filled if _is_number(c))
    return numeric <= len(filled) // 4


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _region_to_table(
    sheet: str,
    first_row: int,
    col_offset: int,
    region: list[list[str]],
    caption: str | None,
    table_no: int,
) -> Element:
    if _looks_like_header(region[0]):
        headers, body, header_row = region[0], region[1:], first_row
    else:
        headers = [f"Column {column_letter(col_offset + i + 1)}" for i in range(len(region[0]))]
        body, header_row = region, first_row - 1
    headers = dedupe_headers([h or f"Column {column_letter(col_offset + i + 1)}" for i, h in enumerate(headers)])
    last_row = first_row + len(region) - 1
    last_col = col_offset + len(headers)
    cell_range = f"{column_letter(col_offset + 1)}{first_row}:{column_letter(last_col)}{last_row}"
    return Element(
        kind=ElementKind.TABLE,
        section_path=[sheet],
        table=TableData(
            headers=headers,
            rows=body,
            caption=caption,
            origin_row=header_row,
            origin_col=col_offset + 1,
        ),
        location=SourceLocation(sheet=sheet, cell_range=cell_range),
        attrs={"table_index": table_no},
    )
