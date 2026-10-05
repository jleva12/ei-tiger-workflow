"""PowerPoint (.pptx) -> IR using python-pptx (pure Python, no ML models).

Each visible slide is a section named by its title ("Slide N" when it has
none): the title becomes a level-1 heading, and the slide's shapes follow in
reading order (rows top to bottom, left to right within a row; group shapes
flattened in place):

  text      body placeholders are bullets unless a paragraph turns them off;
            text boxes are paragraphs unless one turns them on. Indent levels
            are kept. Date, footer and slide-number placeholders are skipped.
  tables    TABLE elements, spanned cells blanked so columns stay aligned.
  charts    their data as a TABLE (categories x series), captioned with the
            chart's title.
  SmartArt  its nodes' text, as list items.
  notes     the speaker notes, under "<slide title> / Notes".

Hidden slides are skipped. Every element carries its slide number (``page``)
and title (``page_name``). Slide shows (.ppsx) and macro-enabled decks
(.pptm, .ppsm) read the same way; legacy binary .ppt does not (save it as
.pptx).
"""

from __future__ import annotations

import io
import zipfile
from typing import Any

from forge_task_documents.errors import ParseError
from forge_task_documents.models import (
    Element,
    ElementKind,
    ParsedDocument,
    SourceFile,
    SourceLocation,
    TableData,
)
from forge_task_documents.parsers.base import (
    SectionTracker,
    cell_to_str,
    clean_text,
    core_properties_metadata,
    dedupe_headers,
    title_from_filename,
)

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

# python-pptx opens presentations only; a slide show is the same package under another main type.
_MAIN_TYPES = {
    b"application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml": (
        b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    ),
    b"application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml": (
        b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"
    ),
}

# Shapes shown in a row are read left to right: tops within this (EMU, 0.25") share a row.
_ROW_TOLERANCE = 228_600


class PptxParser:
    name = "pptx"
    version = "1"
    extensions = frozenset({"pptx", "pptm", "ppsx", "ppsm"})
    media_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/vnd.openxmlformats-officedocument.presentationml.slideshow",
            "application/vnd.ms-powerpoint.presentation.macroenabled.12",
            "application/vnd.ms-powerpoint.slideshow.macroenabled.12",
        }
    )

    def __init__(self, *, include_hidden_slides: bool = False, include_notes: bool = True) -> None:
        self.include_hidden_slides = include_hidden_slides
        self.include_notes = include_notes

    def parse(self, source: SourceFile) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise ParseError("python-pptx is not installed") from exc

        try:
            prs = Presentation(io.BytesIO(_as_presentation(source.data)))
        except Exception as exc:
            raise ParseError(f"cannot open pptx {source.filename!r}: {exc}") from exc

        tracker = SectionTracker()
        elements: list[Element] = []
        slides: list[str] = []
        deck_title = ""
        try:
            for number, slide in enumerate(prs.slides, start=1):
                if not self.include_hidden_slides and slide._element.get("show") in ("0", "false"):
                    continue
                title_shape = slide.shapes.title
                title = _one_line(title_shape.text_frame.text) if title_shape is not None else ""
                name = title or f"Slide {number}"
                if number == 1 and title and _is_title_slide(title_shape):
                    deck_title = title
                slides.append(name)
                location = SourceLocation(page=number, page_name=name)
                path = tracker.push(1, name)
                elements.append(
                    Element(kind=ElementKind.HEADING, text=name, level=1, section_path=path, location=location)
                )
                skip = title_shape.shape_id if title_shape is not None else None
                for shape in _reading_order(slide.shapes):
                    if shape.shape_id != skip:
                        elements.extend(_shape_elements(shape, slide, path, location))
                if self.include_notes:
                    elements.extend(_notes(slide, [*path, "Notes"], location))
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError(f"cannot read pptx {source.filename!r}: {exc}") from exc

        core = prs.core_properties
        title = (core.title or "").strip() or deck_title or title_from_filename(source.filename)
        return ParsedDocument(
            tenant_id=source.tenant_id,
            doc_id=source.doc_id,
            source_type="pptx",
            title=title,
            elements=elements,
            parser_name=self.name,
            parser_version=self.version,
            metadata={**core_properties_metadata(core), "slides": slides},
        )


def _as_presentation(data: bytes) -> bytes:
    """The file with a slide show's main part retyped as a presentation's."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zin:
            types = zin.read("[Content_Types].xml")
            fixed = types
            for show, presentation in _MAIN_TYPES.items():
                fixed = fixed.replace(show, presentation)
            if fixed == types:
                return data
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    zout.writestr(item, fixed if item.filename == "[Content_Types].xml" else zin.read(item.filename))
            return out.getvalue()
    except (zipfile.BadZipFile, KeyError):
        return data  # let python-pptx say what's wrong with it


def _one_line(text: str) -> str:
    return clean_text(text.replace("\v", "\n")).replace("\n", " ")


def _is_title_slide(shape: Any) -> bool:
    from pptx.enum.shapes import PP_PLACEHOLDER

    try:
        return bool(shape.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE)
    except ValueError:
        return False


def _reading_order(shapes: Any) -> list[Any]:
    """Rows of shapes top to bottom, each row left to right."""
    placed = sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))
    rows: list[list[Any]] = []
    for shape in placed:
        if rows and (shape.top or 0) - (rows[-1][0].top or 0) <= _ROW_TOLERANCE:
            rows[-1].append(shape)
        else:
            rows.append([shape])
    return [shape for row in rows for shape in sorted(row, key=lambda s: s.left or 0)]


def _shape_elements(shape: Any, slide: Any, path: list[str], location: SourceLocation) -> list[Element]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [el for child in _reading_order(shape.shapes) for el in _shape_elements(child, slide, path, location)]
    if getattr(shape, "has_table", False):
        table = _table(shape.table)
        return [Element(kind=ElementKind.TABLE, table=table, section_path=path, location=location)] if table else []
    if getattr(shape, "has_chart", False):
        table = _chart_table(shape.chart)
        return [Element(kind=ElementKind.TABLE, table=table, section_path=path, location=location)] if table else []
    if _is_smartart(shape):
        return [
            Element(kind=ElementKind.LIST_ITEM, text=text, level=0, section_path=path, location=location)
            for text in _smartart_texts(shape, slide)
        ]
    if not getattr(shape, "has_text_frame", False):
        return []

    bulleted = False
    if shape.is_placeholder:
        try:
            kind = shape.placeholder_format.type
        except ValueError:
            kind = None
        if kind in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER):
            return []
        bulleted = kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)

    out: list[Element] = []
    for paragraph in shape.text_frame.paragraphs:
        text = clean_text(paragraph.text.replace("\v", "\n"))
        if not text:
            continue
        if _bullet(paragraph, default=bulleted):
            out.append(
                Element(
                    kind=ElementKind.LIST_ITEM,
                    text=text,
                    level=paragraph.level,
                    section_path=path,
                    location=location,
                )
            )
        else:
            out.append(Element(kind=ElementKind.PARAGRAPH, text=text, section_path=path, location=location))
    return out


def _bullet(paragraph: Any, *, default: bool) -> bool:
    """Whether a paragraph shows a bullet: its own setting, else its shape's default."""
    ppr = paragraph._p.pPr
    if ppr is not None:
        for child in ppr:
            tag = str(child.tag)
            if tag == f"{_A}buNone":
                return False
            if tag in (f"{_A}buChar", f"{_A}buAutoNum", f"{_A}buBlip"):
                return True
    return default


def _table(table: Any) -> TableData | None:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = ["" if cell.is_spanned else _one_line(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    if table.first_row and len(rows) > 1:
        header, body = rows[0], rows[1:]
    else:
        header, body = [""] * width, rows
    return TableData(headers=dedupe_headers([h or f"Column {i + 1}" for i, h in enumerate(header)]), rows=body)


def _chart_table(chart: Any) -> TableData | None:
    try:
        plots = list(chart.plots)
        if not plots:
            return None
        categories = [cell_to_str(c) for c in plots[0].categories]
        series = [
            (clean_text(s.name or "") or f"Series {i}", list(s.values))
            for i, s in enumerate((s for plot in plots for s in plot.series), start=1)
        ]
    except Exception:  # a chart python-pptx can't read (external data, chartex): skip it
        return None
    if not series:
        return None
    length = max([len(categories), *(len(values) for _, values in series)])
    if not categories:
        categories = [str(i) for i in range(1, length + 1)]
    rows = [
        [categories[i] if i < len(categories) else "", *(cell_to_str(v[i]) if i < len(v) else "" for _, v in series)]
        for i in range(length)
    ]
    title = ""
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = _one_line(chart.chart_title.text_frame.text)
    except Exception:
        pass
    return TableData(
        headers=dedupe_headers(["Category", *(name for name, _ in series)]),
        rows=rows,
        caption=title or "Chart",
    )


def _is_smartart(shape: Any) -> bool:
    data = shape._element.find(f".//{_A}graphicData")
    return data is not None and data.get("uri") == _DIAGRAM_URI


def _smartart_texts(shape: Any, slide: Any) -> list[str]:
    ids = shape._element.find(f".//{_DGM}relIds")
    if ids is None:
        return []
    try:
        part = slide.part.related_part(ids.get(f"{_R}dm"))
        return diagram_texts(part.blob)
    except Exception:  # a missing or unreadable data part: nothing to add
        return []


def diagram_texts(blob: bytes) -> list[str]:
    """The text of each node in a SmartArt data part, in document order."""
    from lxml import etree

    root = etree.fromstring(blob)
    texts: list[str] = []
    for point in root.iter(f"{_DGM}pt"):
        if point.get("type", "node") != "node":
            continue  # doc, pres, parTrans, sibTrans: structure, not content
        body = point.find(f"{_DGM}t")
        if body is None:
            continue
        paragraphs = ["".join(t.text or "" for t in p.iter(f"{_A}t")) for p in body.iter(f"{_A}p")]
        text = clean_text(" ".join(p for p in paragraphs if p.strip()))
        if text:
            texts.append(text)
    return texts


def _notes(slide: Any, path: list[str], location: SourceLocation) -> list[Element]:
    if not slide.has_notes_slide:
        return []
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return []
    out: list[Element] = []
    for paragraph in frame.paragraphs:
        text = clean_text(paragraph.text.replace("\v", "\n"))
        if text:
            out.append(Element(kind=ElementKind.PARAGRAPH, text=text, section_path=path, location=location))
    return out
